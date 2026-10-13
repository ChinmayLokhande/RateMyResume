import io

import pymupdf
import pytest
from docx import Document

from app.services.extraction import ExtractionError, extract_docx, extract_pdf

LONG_LINE = "Test Person - test@example.com - Python, Node.js, PostgreSQL developer. " * 3


def make_pdf(text: str | None, password: str | None = None) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    if text:
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 750), text)
    if password:
        return doc.tobytes(
            encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password + "x"
        )
    return doc.tobytes()


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Test Person")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Python"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_with_text_is_extracted_natively():
    result = extract_pdf(make_pdf(LONG_LINE))
    assert "test@example.com" in result.text
    assert result.needs_ocr is False
    assert result.page_count == 1


def test_blank_pdf_is_flagged_for_ocr():
    result = extract_pdf(make_pdf(None))
    assert result.needs_ocr is True


def test_docx_paragraphs_and_tables_are_extracted():
    result = extract_docx(make_docx())
    assert "Test Person" in result.text
    assert "Skills | Python" in result.text


def test_password_protected_pdf_is_rejected():
    with pytest.raises(ExtractionError, match="password"):
        extract_pdf(make_pdf(LONG_LINE, password="secret"))


def test_garbage_bytes_are_rejected():
    with pytest.raises(ExtractionError):
        extract_pdf(b"not a pdf at all")
    with pytest.raises(ExtractionError):
        extract_docx(b"not a docx at all")